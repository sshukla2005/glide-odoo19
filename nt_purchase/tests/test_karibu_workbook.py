import datetime
import io

from openpyxl import load_workbook

from odoo.tests import BaseCase, tagged

from odoo.addons.nt_purchase.tools.karibu_workbook import build_karibu_workbook

PRODUCTS = [
    {'code': 'BAG325N', 'description': 'CEMENT;32.5N;50KG-BAGS;(SIMBA BARABARA)',
     'short_type': '32.5N;50KG'},
    {'code': 'BAG325R', 'description': 'CEMENT;32.5R;50KG-BAGS;(SIMBA IMARA)',
     'short_type': '32.5R;50KG'},
]


def _lpo(lines):
    return {
        'reference': 'GPAL-10872-01',
        'date': datetime.date(2026, 6, 3),
        'location': 'IRINGA',
        'customer_code': 'C60010467',
        'signatory': 'SUNIL KUMAR',
        'company': {'name': 'Glide Paths Africa Limited', 'address': 'Masaki, Dar es Salaam',
                    'phone': '+255767100535', 'email': 'sales@glidepathsafrica.co.tz',
                    'tin': '158-550-024', 'vrn': '40-050592-E'},
        'supplier': {'name': 'Tanga Cement Public Limited Company', 'street': 'Pongwe',
                     'street2': 'Tanga', 'phone': '', 'email': ''},
        'lines': lines,
    }


def _line(**overrides):
    line = {'driver': 'RAMADHANI MOSHI', 'license': '4000539401', 'truck': 'T740EHQ',
            'trailer': 'T123ABC', 'qty': 31.0, 'product_code': 'BAG325R',
            'product_name': 'CEMENT;32.5R;50KG-BAGS;(SIMBA IMARA)',
            'product_type': '32.5R;50KG', 'transporter': 'Glide Paths Africa Limited'}
    line.update(overrides)
    return line


@tagged('post_install', '-at_install')
class TestKaribuWorkbook(BaseCase):

    def _reload(self, wb, data_only=False):
        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        return load_workbook(stream, data_only=data_only)

    def test_sheets_and_bulk_upload(self):
        wb = self._reload(build_karibu_workbook(
            _lpo([_line(), _line(truck='T275EQX', license='DL-77', qty=30.5)]), PRODUCTS))
        self.assertEqual(wb.sheetnames, ['Products', 'BulkUpload', 'Sheet1'])
        bulk = wb['BulkUpload']
        self.assertEqual(
            [c.value for c in bulk[5]][:7],
            ['No', 'Order reference no.', 'Product', 'Product Name',
             'Truck Reg no.', 'Driver Licence no.', 'Qnty (tons)'])
        self.assertEqual(
            [c.value for c in bulk[6]][:7],
            [1, 'GPAL-10872-01', 'CEMENT;32.5R;50KG-BAGS;(SIMBA IMARA)', 'BAG325R',
             'T740EHQ', 4000539401, 31])
        # Non-numeric licences stay text.
        self.assertEqual(bulk['F7'].value, 'DL-77')
        # Empty rows keep the client's lookup formula for manual entry.
        self.assertTrue(str(bulk['C8'].value).startswith('=IF(TRIM(D8)'))
        self.assertEqual(bulk['A106'].value, 101)

    def test_product_readable_without_excel(self):
        """A parser that reads values only must see the product description."""
        wb = self._reload(build_karibu_workbook(_lpo([_line()]), PRODUCTS), data_only=True)
        self.assertEqual(wb['BulkUpload']['C6'].value, 'CEMENT;32.5R;50KG-BAGS;(SIMBA IMARA)')

    def test_products_sheet(self):
        wb = self._reload(build_karibu_workbook(_lpo([_line()]), PRODUCTS))
        products = wb['Products']
        self.assertEqual(products['A2'].value, 'BAG325N')
        self.assertEqual(products['B3'].value, 'CEMENT;32.5R;50KG-BAGS;(SIMBA IMARA)')

    def test_printed_lpo_total(self):
        lines = [_line(qty=31.0) for _i in range(25)]
        wb = self._reload(build_karibu_workbook(_lpo(lines), PRODUCTS))
        sheet = wb['Sheet1']
        self.assertEqual(sheet['I12'].value, 'GPAL-10872-01')
        self.assertEqual(sheet['G11'].value, 'C60010467')
        # 25 trucks from row 19: the total sums all of them, in the MT column.
        self.assertEqual(sheet['H44'].value, '=SUM(H19:H43)')
        self.assertEqual(sheet['D44'].value, 'GRAND TOTAL:::')
